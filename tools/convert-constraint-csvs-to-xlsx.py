#!/usr/bin/env python3
"""Convert legacy ClearPlan CSV catalogs into one normalized XLSX workbook."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TABLE_HEADERS = [
    "table_id",
    "display_name",
    "active",
    "is_plan_sum",
    "fx_min",
    "fx_max",
    "dpf_min_gy",
    "dpf_max_gy",
    "total_dose_min_gy",
    "total_dose_max_gy",
    "site",
    "regime",
    "source",
]

CONSTRAINT_HEADERS = [
    "constraint_id",
    "table_id",
    "structure_id",
    "structure_name",
    "metric",
    "unit",
    "comparator",
    "goal",
    "variation",
    "priority",
    "source",
    "comment",
    "legacy_metadata",
]

STRUCTURE_HEADERS = [
    "structure_id",
    "canonical_name",
    "active",
    "laterality",
    "aliases",
    "side_aliases_left",
    "side_aliases_right",
    "dicom_type",
    "codes",
]

FIXED_TIME = datetime(2026, 7, 28, 8, 0, 0, tzinfo=timezone.utc)


def read_csv(path: Path) -> list[dict[str, str]]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return [
                    {str(key).strip(): (value or "").strip() for key, value in row.items()}
                    for row in csv.DictReader(handle)
                ]
        except UnicodeDecodeError as error:
            last_error = error
    raise ValueError(f"Cannot decode {path}: {last_error}")


def split_values(value: str) -> list[str]:
    return list(
        OrderedDict.fromkeys(
            item.strip()
            for item in (value or "").split("|")
            if item and item.strip()
        )
    )


def parse_number(value: Any) -> float | int | None:
    text = "" if value is None else str(value).strip()
    if not text:
        return None
    text = text.replace(",", ".")
    number = float(text)
    return int(number) if number.is_integer() else number


def normalize_comparator(value: str) -> str:
    return (value or "").strip().replace("≤", "<=").replace("≥", ">=")


def parse_evaluation(value: str) -> tuple[str, float | int | None]:
    match = re.match(
        r"^\s*(<=|>=|<|>|≤|≥)?\s*(-?\d+(?:[.,]\d+)?)?\s*$",
        value or "",
    )
    if not match:
        raise ValueError(f"Unsupported evaluation point: {value!r}")
    comparator = normalize_comparator(match.group(1) or "")
    return comparator, parse_number(match.group(2))


def normalize_metric(metric: str) -> str:
    compact = re.sub(r"\s+", "", metric or "").replace(",", ".")
    lower = compact.casefold()
    if lower in {"mean", "dmean"}:
        return "Dmean"
    if lower in {"max", "dmax"}:
        return "Dmax"
    if lower in {"min", "dmin"}:
        return "Dmin"
    if lower.startswith("cv"):
        return "CV" + compact[2:]
    if lower.startswith("d"):
        return "D" + compact[1:]
    if lower.startswith("v"):
        return "V" + compact[1:]
    return compact


def normalize_unit(unit: str) -> str:
    stripped = (unit or "").strip()
    if stripped.casefold() in {"cm³", "cc"}:
        return "cc"
    if stripped.casefold() == "gy":
        return "Gy"
    if stripped.casefold().startswith("cm ("):
        return "cm"
    return stripped


def parse_objective(value: str) -> tuple[str, str]:
    match = re.match(r"^\s*(.*?)\s*(?:\[([^\]]*)\])?\s*$", value or "")
    if not match:
        raise ValueError(f"Unsupported DVH objective: {value!r}")
    return normalize_metric(match.group(1)), normalize_unit(match.group(2) or "")


def parse_priority(value: str) -> int | None:
    match = re.match(r"^\s*(\d+)", value or "")
    return int(match.group(1)) if match else None


def table_metadata(path: Path) -> dict[str, Any]:
    table_id = path.stem
    name_lower = table_id.casefold()
    match = re.search(r"(?:^|_)(\d+)fx$", name_lower)
    fx = int(match.group(1)) if match else None
    if "hypofractionated" in name_lower:
        fx_min, fx_max = 1, 5
    else:
        fx_min = fx_max = fx
    return {
        "table_id": table_id,
        "display_name": table_id.replace("_", " "),
        "active": True,
        "is_plan_sum": "sum" in name_lower,
        "fx_min": fx_min,
        "fx_max": fx_max,
        "dpf_min_gy": None,
        "dpf_max_gy": None,
        "total_dose_min_gy": None,
        "total_dose_max_gy": None,
        "site": "",
        "regime": "Conventional" if "conv" in name_lower else "",
        "source": f"Legacy migration from {path.name}",
    }


def canonical_legacy_row(row: dict[str, str]) -> dict[str, Any]:
    if "TemplateId" in row:
        return {
            "structure_id": row.get("TemplateId", ""),
            "structure_name": row.get("TemplateId", ""),
            "codes": row.get("TemplateCodes", ""),
            "aliases": row.get("TemplateAliases", ""),
            "dicom_type": row.get("TemplateTypes", ""),
            "objective": row.get("DVHObjective", ""),
            "evaluation": row.get("Goal", ""),
            "variation": row.get("Variation", ""),
            "priority": row.get("Priority", ""),
            "source": row.get("Source", ""),
            "comment": row.get("Notes", ""),
        }
    return {
        "structure_id": row.get("Structure IDs", ""),
        "structure_name": row.get("Structure IDs", ""),
        "codes": row.get("Structure Codes", ""),
        "aliases": row.get("IDAliases", ""),
        "dicom_type": row.get("CodeAliases", ""),
        "objective": row.get("DVH Objective", ""),
        "evaluation": row.get("Evaluation Point", ""),
        "variation": row.get("Variation", ""),
        "priority": row.get("Priority", ""),
        "source": row.get("Source", ""),
        "comment": row.get("ZusatzInfo", ""),
    }


def infer_laterality(structure_id: str) -> str:
    normalized = (structure_id or "").replace(" ", "")
    if "R+L" in normalized:
        return "R+L"
    if "L/R" in normalized:
        return "L/R"
    if normalized.endswith(("_L", "-L")):
        return "L"
    if normalized.endswith(("_R", "-R")):
        return "R"
    return ""


def convert(inputs: list[Path]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    tables: list[dict[str, Any]] = []
    constraints: list[dict[str, Any]] = []
    structures: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for input_path in inputs:
        table = table_metadata(input_path)
        tables.append(table)
        for row_number, raw_row in enumerate(read_csv(input_path), start=1):
            legacy = canonical_legacy_row(raw_row)
            metric, unit = parse_objective(legacy["objective"])
            comparator, goal = parse_evaluation(legacy["evaluation"])
            _, variation = parse_evaluation(legacy["variation"])
            structure_id = legacy["structure_id"] or legacy["structure_name"]
            structure_key = structure_id.strip().casefold()
            if structure_key not in structures:
                structures[structure_key] = {
                    "structure_id": structure_id,
                    "canonical_name": legacy["structure_name"] or structure_id,
                    "active": True,
                    "laterality": infer_laterality(structure_id),
                    "aliases": set(),
                    "side_aliases_left": set(),
                    "side_aliases_right": set(),
                    "dicom_type": set(),
                    "codes": set(),
                }
            structure = structures[structure_key]
            canonical_structure_id = structure["structure_id"]
            constraint = {
                "constraint_id": f"{table['table_id']}-{row_number:04d}",
                "table_id": table["table_id"],
                "structure_id": canonical_structure_id,
                "structure_name": legacy["structure_name"],
                "metric": metric,
                "unit": unit,
                "comparator": comparator,
                "goal": goal,
                "variation": variation,
                "priority": parse_priority(legacy["priority"]),
                "source": legacy["source"],
                "comment": legacy["comment"],
                "legacy_metadata": json.dumps(
                    raw_row, ensure_ascii=False, sort_keys=True, separators=(",", ":")
                ),
            }
            constraints.append(constraint)
            structure["aliases"].update(split_values(legacy["aliases"]))
            structure["dicom_type"].update(split_values(legacy["dicom_type"]))
            structure["codes"].update(split_values(legacy["codes"]))
    normalized_structures: list[dict[str, Any]] = []
    for structure in structures.values():
        normalized_structures.append(
            {
                key: (
                    "|".join(sorted(value, key=str.casefold))
                    if isinstance(value, set)
                    else value
                )
                for key, value in structure.items()
            }
        )
    return tables, constraints, normalized_structures


def write_rows(worksheet, headers: list[str], rows: Iterable[dict[str, Any]]) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315A8A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column_index, header in enumerate(headers, start=1):
        maximum = max(
            len(str(worksheet.cell(row=row_index, column=column_index).value or ""))
            for row_index in range(1, worksheet.max_row + 1)
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(maximum + 2, len(header) + 2), 48
        )


def build_workbook(
    tables: list[dict[str, Any]],
    constraints: list[dict[str, Any]],
    structures: list[dict[str, Any]],
) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "ClearPlan"
    workbook.properties.lastModifiedBy = "ClearPlan"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    write_rows(workbook.create_sheet("Tables"), TABLE_HEADERS, tables)
    write_rows(
        workbook.create_sheet("Constraints"), CONSTRAINT_HEADERS, constraints
    )
    write_rows(workbook.create_sheet("Structures"), STRUCTURE_HEADERS, structures)
    readme = workbook.create_sheet("README")
    readme["A1"] = "ClearPlan Constraint Catalog"
    readme["A1"].font = Font(bold=True, size=16, color="315A8A")
    readme["A3"] = "Schema"
    readme["B3"] = "ClearPlan.constraint_catalog.xlsx.v1"
    readme["A5"] = (
        "Die Blätter Tables, Constraints und Structures bilden die interne "
        "Zuordnung ab. Aliase werden mit | getrennt."
    )
    readme["A7"] = "Unicode-Prüfung: ä ö ü Ä Ö Ü ß ≤ ≥ cm³ Größe Rückenmark."
    readme.column_dimensions["A"].width = 92
    readme.column_dimensions["B"].width = 42
    return workbook


def workbook_signature(workbook: Workbook) -> tuple[Any, ...]:
    return tuple(
        (
            worksheet.title,
            tuple(
                tuple(
                    None if cell.value in (None, "") else cell.value
                    for cell in row
                )
                for row in worksheet.iter_rows()
            ),
        )
        for worksheet in workbook.worksheets
    )


def write_and_verify(
    output: Path,
    tables: list[dict[str, Any]],
    constraints: list[dict[str, Any]],
    structures: list[dict[str, Any]],
) -> None:
    workbook = build_workbook(tables, constraints, structures)
    expected_signature = workbook_signature(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    reopened = load_workbook(output, read_only=True, data_only=True)
    actual_signature = workbook_signature(reopened)
    reopened.close()
    if actual_signature != expected_signature:
        raise RuntimeError("Workbook content changed during XLSX round-trip.")
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        second_path = Path(handle.name)
    try:
        second = build_workbook(tables, constraints, structures)
        second.save(second_path)
        reopened_second = load_workbook(second_path, read_only=True, data_only=True)
        second_signature = workbook_signature(reopened_second)
        reopened_second.close()
        if second_signature != expected_signature:
            raise RuntimeError("Second workbook build is not cell-deterministic.")
    finally:
        second_path.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", action="append", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    inputs = [path.resolve() for path in args.input]
    for path in inputs:
        if not path.is_file():
            parser.error(f"Input CSV does not exist: {path}")
    tables, constraints, structures = convert(inputs)
    write_and_verify(args.output.resolve(), tables, constraints, structures)
    print(
        f"tables={len(tables)} constraints={len(constraints)} "
        f"structures={len(structures)} output={args.output.resolve()}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
