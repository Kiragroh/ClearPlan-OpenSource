"""Read the authorized 2024 source workbook; emit JSON for the XLSX authoring step.

Never writes an XLSX, modifies source cells, recalculates cached formulas, or
infers patient-specific applicability. All semantic conversions live in JSON.
"""
from __future__ import annotations

import argparse
import copy
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

TABLE_HEADERS = ["table_id", "display_name", "active", "is_plan_sum", "fx_min", "fx_max", "dpf_min_gy", "dpf_max_gy", "total_dose_min_gy", "total_dose_max_gy", "site", "regime", "source", "requires_confirmation"]
CONSTRAINT_HEADERS = ["constraint_id", "table_id", "structure_id", "structure_name", "metric", "unit", "comparator", "goal", "variation", "priority", "source", "comment", "source_sheet", "source_row", "reference_codes", "raw_objective", "raw_goal", "raw_variation"]
STRUCTURE_HEADERS = ["structure_id", "canonical_name", "active", "laterality", "aliases", "side_aliases_left", "side_aliases_right", "dicom_type", "codes", "source_ids", "naming_source"]
SOURCE_HEADERS = ["source_id", "source_label", "original_code", "reference_ids", "reference_text", "url", "source_location", "verification"]
REVIEW_HEADERS = ["record_id", "kind", "table_id", "source_sheet", "source_row", "source_structure", "canonical_candidate", "raw_objective", "raw_goal", "raw_variation", "reference_codes", "context", "reason"]
METRIC = re.compile(r"^(Mean|Max|Min|Volume|D\d+(?:\.\d+)?(?:%|cc)|(?:V|CV)\d+(?:\.\d+)?(?:Gy|%))\[(Gy|cGy|%|cc|cm3)\]$")
GOAL = re.compile(r"^(<=|>=|<|>)\s*([0-9]+(?:\.[0-9]+)?)$")


def text(value):
    return "" if value is None else str(value).strip()


def key(value):
    return unicodedata.normalize("NFKC", text(value)).casefold()


def match_key(value):
    """Mirror StructureAliasResolver.NormalizeName, including TG263 qualifiers."""
    return "".join(c for c in key(value) if c.isalnum() or c in "~-+^!=/")


def number(value):
    if value is None or text(value) == "":
        return None
    try:
        parsed = Decimal(text(value))
    except InvalidOperation as exc:
        raise ValueError("Not a scalar number") from exc
    if not parsed.is_finite():
        raise ValueError("Not a finite scalar number")
    return int(parsed) if parsed == parsed.to_integral() else float(parsed)


def parse_objective(objective, goal, variation):
    metric_match = METRIC.fullmatch(text(objective))
    goal_match = GOAL.fullmatch(text(goal))
    if not metric_match:
        raise ValueError("Unsupported or malformed DVH objective; no silent syntax repair")
    if not goal_match:
        raise ValueError("Goal is not a single explicit comparator and scalar")
    metric, unit = metric_match.groups()
    comparator, raw_goal = goal_match.groups()
    goal_value, variation_value = number(raw_goal), number(variation)
    if variation_value is not None:
        if comparator.startswith("<") and variation_value < goal_value:
            raise ValueError("Variation is stricter than goal for an upper-bound objective")
        if comparator.startswith(">") and variation_value > goal_value:
            raise ValueError("Variation is stricter than goal for a lower-bound objective")
    if metric.startswith("CV") and not comparator.startswith(">"):
        raise ValueError("Complementary-volume metric has a non-lower-bound comparator")
    return {"metric": {"Mean": "Dmean", "Max": "Dmax", "Min": "Dmin"}.get(metric, metric), "unit": unit, "comparator": comparator, "goal": goal_value, "variation": variation_value}


def drawing_references(path):
    """The real source-code legend is in DrawingML, not worksheet cells."""
    references, legend = {}, []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith("xl/drawings/drawing") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            for index, element in enumerate(root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"), 1):
                value = text(element.text)
                match = re.match(r"^(\d+)\.\s+(.*)$", value, re.DOTALL)
                if match:
                    references.setdefault(match[1], {"text": match[2], "location": name + " text " + str(index)})
                elif value and ("stands for" in value or "C_Letter_Number" in value or "international guidelines" in value):
                    if value not in [item["text"] for item in legend]:
                        legend.append({"text": value, "location": name + " text " + str(index)})
    return references, legend


def source_record(code, references):
    original = text(code)
    ids = []
    if re.search(r"(?:^|[ ,_])T(?:$|[ ,_(\-])", original):
        ids.append("1")
    if original.startswith("C_"):
        ids.append("2")
        suffix = original.split("_", 2)[-1]
        if re.fullmatch(r"\d+(?:,\d+)*", suffix):
            ids.extend(suffix.split(","))
    elif re.fullmatch(r"\d+", original):
        ids.append(original)
    elif original == "HIPPORAD, 3":
        ids.append("3")
    ids = list(dict.fromkeys(ids))
    label = original.replace("UKE", "Institutional 2024")
    if original == "T":
        label = "Timmerman 2022 (2024 compilation)"
    elif original.startswith("C_"):
        label = "CORSAIR " + original
    descriptions = [f"[{n}] {references[n]['text']}" for n in ids if n in references]
    missing = [n for n in ids if n not in references]
    verification = "Copied from the source workbook legend; numeric constraints have not been revalidated against primary literature."
    if missing:
        verification += " References absent from workbook legend: " + ", ".join(missing) + "."
    if original.startswith("C_") and "-" in original:
        verification += " Ambiguous reference range retained without expansion."
    urls = []
    if "1" in ids:
        urls.append("https://doi.org/10.1016/j.ijrobp.2021.09.027")
    if "2" in ids:
        urls.append("https://doi.org/10.3390/curroncol29100552")
    return {"source_id": "source-" + hashlib.sha256(original.encode("utf-8")).hexdigest()[:12], "source_label": label or "Source not stated", "original_code": original, "reference_ids": "|".join(ids), "reference_text": "\n".join(descriptions), "url": "|".join(urls), "source_location": "; ".join(references[n]["location"] for n in ids if n in references), "verification": verification}


def extract(path, mapping):
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    if before != mapping["source_sha256"]:
        raise ValueError("Source workbook hash differs from reviewed snapshot. Review changes and update mapping explicitly.")
    cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    canonical = {key(k): v for k, v in mapping["canonical_map"].items()}
    unresolved = {key(k): v for k, v in mapping["unresolved_structures"].items()}
    references, legend = drawing_references(path)
    review, sources, constraints, tables = [], {}, [], []
    rejected_aliases = set()
    aliases = defaultdict(set)
    structure_ids = defaultdict(set)
    patterns = [re.compile(p, re.IGNORECASE) for p in mapping["excluded_alias_patterns"]]
    excluded_aliases = {key(k): {key(v) for v in values} for k, values in mapping["excluded_aliases"].items()}

    def add_aliases(source_id, alias_text, location, row_number, sheet):
        canon = canonical.get(key(source_id))
        if not canon:
            return
        structure_ids[canon].add(source_id)
        for alias in [source_id] + text(alias_text).split("|"):
            alias = text(alias)
            if not alias:
                continue
            # PRV identifiers are kept only on explicitly PRV canonical structures.
            bad_pattern = any(p.search(alias) for p in patterns)
            if "PRV" in canon and re.search(r"\bPRV\b", alias, re.IGNORECASE):
                bad_pattern = any(p.search(alias) for p in patterns if "PTV|CTV" not in p.pattern)
            if key(alias) in excluded_aliases.get(key(source_id), set()) or bad_pattern:
                rejection_key = (canon, match_key(alias))
                if rejection_key not in rejected_aliases:
                    rejected_aliases.add(rejection_key)
                    review.append({"record_id": "alias-" + hashlib.sha256((canon + alias).encode("utf-8")).hexdigest()[:12], "kind": "Alias", "table_id": "", "source_sheet": sheet, "source_row": row_number, "source_structure": source_id, "canonical_candidate": canon, "raw_objective": "", "raw_goal": "", "raw_variation": "", "reference_codes": "", "context": alias, "reason": "Alias denotes a different, cropped, combined, ambiguous or placeholder volume; excluded from matching. First source occurrence shown."})
            else:
                aliases[canon].add(alias)

    for row_number, row in enumerate(cached["Glossar"].iter_rows(values_only=True), 1):
        if row_number > 1 and row[0]:
            add_aliases(text(row[0]), row[2], f"Glossar!{row_number}", row_number, "Glossar")

    input_count, formula_count = 0, 0
    for table_config in mapping["tables"]:
        sheet = table_config["sheet"]
        table = {h: None for h in TABLE_HEADERS}
        table.update({k: table_config[k] for k in ("table_id", "display_name", "fx_min", "fx_max", "requires_confirmation")})
        table.update(active=True, is_plan_sum=False, site="", regime="2024 source compilation; local review required", source="Stock 2024 compilation: " + sheet)
        tables.append(table)
        raw_rows = list(formulas[sheet].iter_rows(values_only=True))
        for row_number, row in enumerate(cached[sheet].iter_rows(values_only=True), 1):
            if row_number == 1 or not row[0]:
                continue
            input_count += 1
            source_id, objective, goal, variation = text(row[0]), row[4], row[5], row[6]
            canon = canonical.get(key(source_id), "")
            context, code = text(row[11]), text(row[10])
            location = f"{sheet}!{row_number}"
            original = {"record_id": f"{table['table_id']}-r{row_number}", "kind": "Constraint", "table_id": table["table_id"], "source_sheet": sheet, "source_row": row_number, "source_structure": source_id, "canonical_candidate": canon, "raw_objective": text(objective), "raw_goal": text(goal), "raw_variation": text(variation), "reference_codes": code, "context": context}
            reason = None
            if key(source_id) in unresolved:
                reason = unresolved[key(source_id)]
            elif not canon:
                reason = "No reviewed canonical mapping."
            elif code == "MG":
                reason = "Generic target/body template rule, placeholder tolerance or unsupported index; requires prescription-specific local configuration."
            elif context not in mapping["allowed_context_notes"] and context:
                reason = "Special anatomy, laterality, disease or treatment context requires a separate reviewed rule: " + context
            elif not code:
                reason = "No source code supplied."
            fields = None
            try:
                fields = parse_objective(objective, goal, variation)
            except ValueError as exc:
                reason = reason or str(exc)
            for column in (4, 5, 6):
                if isinstance(raw_rows[row_number - 1][column], str) and raw_rows[row_number - 1][column].startswith("="):
                    formula_count += 1
                    if row[column] is None:
                        reason = "Formula has no cached value; source workbook must be recalculated by its maintainer."
            if reason:
                original["reason"] = reason
                review.append(original)
                continue
            add_aliases(source_id, row[2], location, row_number, sheet)
            if code not in sources:
                sources[code] = source_record(code, references)
            constraint = {h: None for h in CONSTRAINT_HEADERS}
            constraint.update(fields)
            constraint.update(constraint_id=original["record_id"], table_id=table["table_id"], structure_id=canon, structure_name=canon, priority=number(row[7]), source=sources[code]["source_label"], comment="; ".join(p for p in [context, "Endpoint: " + text(row[12]) if row[12] else "", "Source row: " + location] if p), source_sheet=sheet, source_row=row_number, reference_codes=code, raw_objective=text(objective), raw_goal=text(goal), raw_variation=text(variation))
            constraints.append(constraint)

    # A shared alias is never silently assigned to whichever structure is listed first.
    owners = defaultdict(set)
    for canon, values in aliases.items():
        for value in values | {canon}:
            owners[match_key(value)].add(canon)
    collisions = {value: sorted(owners[value]) for value in owners if len(owners[value]) > 1}
    structures = []
    for canon in sorted(aliases):
        clean = sorted({value for value in aliases[canon] if match_key(value) not in collisions}, key=key)
        structures.append({"structure_id": canon, "canonical_name": canon, "active": True, "laterality": "left" if canon.endswith("_L") else "right" if canon.endswith("_R") else "", "aliases": "|".join(clean), "side_aliases_left": "", "side_aliases_right": "", "dicom_type": "", "codes": "", "source_ids": "|".join(sorted(structure_ids[canon], key=key)), "naming_source": "AAPM TG263 2017-08-15 primary name"})
    for value, owners_list in collisions.items():
        review.append({"record_id": "collision-" + hashlib.sha256(value.encode("utf-8")).hexdigest()[:12], "kind": "Alias", "table_id": "", "source_sheet": "Glossar and selected tables", "source_row": None, "source_structure": "|".join(owners_list), "canonical_candidate": "", "raw_objective": "", "raw_goal": "", "raw_variation": "", "reference_codes": "", "context": value, "reason": "Normalized alias is shared by multiple canonical structures; excluded from all matching."})
    # Include codes of review-only rows in the provenance ledger as well.
    for item in review:
        code = item["reference_codes"]
        if code and code not in sources:
            sources[code] = source_record(code, references)
    after = hashlib.sha256(path.read_bytes()).hexdigest()
    if after != before:
        raise RuntimeError("Source changed during read; discard extraction")
    review_constraints = sum(item["kind"] == "Constraint" for item in review)
    if input_count != len(constraints) + review_constraints:
        raise AssertionError("Every scoped input constraint must be executable or review-only")
    return {"metadata": {"schema": "ClearPlan.stock_catalog.intermediate.v1", "source_name": path.name, "source_sha256": before, "naming_evidence": mapping["naming_evidence"], "input_constraint_count": input_count, "executable_constraint_count": len(constraints), "review_constraint_count": review_constraints, "review_alias_count": len(review) - review_constraints, "cached_objective_formula_cells": formula_count, "policy_notes": mapping["policy_notes"], "source_legend": legend, "excluded_sheets": mapping["excluded_sheets"], "alias_collisions": collisions, "counts_by_table": dict(Counter(row["table_id"] for row in constraints)), "headers": {"tables": TABLE_HEADERS, "constraints": CONSTRAINT_HEADERS, "structures": STRUCTURE_HEADERS, "reviewRows": REVIEW_HEADERS, "sources": SOURCE_HEADERS}}, "tables": tables, "constraints": constraints, "structures": structures, "reviewRows": review, "sources": list(sources.values())}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--mapping", type=Path, default=Path(__file__).with_name("stock2024.mapping.json"))
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--public-output", type=Path, help="Optional institution-neutral JSON for the distributed XLSX; original codes remain in --output only")
    args = parser.parse_args()
    data = extract(args.source, json.loads(args.mapping.read_text(encoding="utf-8")))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.public_output:
        public = public_catalog(data)
        args.public_output.parent.mkdir(parents=True, exist_ok=True)
        args.public_output.write_text(json.dumps(public, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: data["metadata"][k] for k in ("input_constraint_count", "executable_constraint_count", "review_constraint_count", "review_alias_count", "counts_by_table")}, ensure_ascii=True))


def public_catalog(data):
    """Neutral display codes only. Numeric goals, anatomy and references unchanged."""
    public = copy.deepcopy(data)
    public["metadata"]["source_name"] = "2024 local constraint compilation"
    public["metadata"]["public_source_code_map"] = {"LOCAL2024": "Local institutional constraints in the 2024 compilation; not a primary publication. Original institution code is retained in the private audit only."}
    for item in public["metadata"]["source_legend"]:
        item["text"] = item["text"].replace("UKE", "LOCAL2024")
    for collection in ("constraints", "reviewRows", "sources"):
        for row in public[collection]:
            for field in ("source", "source_label", "reference_codes", "original_code"):
                if isinstance(row.get(field), str):
                    row[field] = row[field].replace("UKE", "LOCAL2024").replace("Institutional 2024", "Local 2024")
    return public


if __name__ == "__main__":
    main()
