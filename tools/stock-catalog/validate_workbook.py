"""Independent, read-only XLSX round-trip check; does not author or save workbooks."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

def validate(workbook_path, public_json):
    data = json.loads(Path(public_json).read_text(encoding="utf-8"))
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    checked = 0
    summaries = []
    for sheet_name, key in [("Tables", "tables"), ("Constraints", "constraints"), ("Structures", "structures"), ("ReviewRows", "reviewRows"), ("Sources", "sources")]:
        sheet = workbook[sheet_name]
        headers = data["metadata"]["headers"][key]
        assert [cell.value for cell in sheet[1]] == headers, sheet_name + " headers"
        assert sheet.max_row == len(data[key]) + 1, sheet_name + " row count"
        assert sheet.freeze_panes == "C2", (sheet_name, "header and identifier freeze", sheet.freeze_panes)
        assert not sheet.sheet_view.showGridLines, sheet_name + " gridlines"
        assert len(sheet.tables) == 1, sheet_name + " table filters"
        for r, source in enumerate(data[key], start=2):
            for c, header in enumerate(headers, start=1):
                cell = sheet.cell(r, c)
                expected = source.get(header)
                actual = cell.value
                if expected == "": expected = None
                if isinstance(expected, str) and expected.startswith("=") and actual == "'" + expected:
                    actual = expected
                assert actual == expected, (sheet_name, cell.coordinate, actual, expected)
                assert cell.data_type not in ("e", "f"), (sheet_name, cell.coordinate, cell.data_type)
                if isinstance(expected, (int, float)) and not isinstance(expected, bool):
                    assert isinstance(actual, (int, float)), (sheet_name, cell.coordinate, "number converted to text")
                checked += 1
        summaries.append({"sheet": sheet_name, "rows": len(data[key]), "freeze": sheet.freeze_panes})
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type not in ("f", "e"), (sheet.title, cell.coordinate, "unexpected formula/error")
                if isinstance(cell.value, str):
                    assert "UKE" not in cell.value and "C:\\Users" not in cell.value, (sheet.title, cell.coordinate, "private label/path")
    workbook.close()
    return {"workbook": str(Path(workbook_path).resolve()), "sha256": hashlib.sha256(Path(workbook_path).read_bytes()).hexdigest(), "checked_cells": checked, "sheets": summaries, "errors": 0}

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("public_json")
    args = parser.parse_args()
    print(json.dumps(validate(args.workbook, args.public_json), indent=2))
