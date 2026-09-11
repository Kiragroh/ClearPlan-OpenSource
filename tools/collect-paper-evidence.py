#!/usr/bin/env python3
"""Collect bounded, patient-free evidence used by the ClearPlan technical note."""

from __future__ import annotations

import hashlib
import argparse
import json
import re
import subprocess
import sys
import unittest
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCENARIO_DIR = ROOT / "ClearPlan.Simulator" / "Scenarios"
MIXED_SCENARIO = SCENARIO_DIR / "mixed-review.json"
WORKBOOK = (
    ROOT
    / "ClearPlan.Script"
    / "Distribution"
    / "ConstraintTemplates"
    / "ClearPlan_DefaultConstraints.xlsx"
)
OUTPUT = ROOT / "paper" / "evidence" / "technical_note_evidence.json"


def parse_csharp_run(output: str, expected_names: list[str] | None = None) -> dict:
    summaries = re.findall(r"(?m)^(\d+)\s+(?:RTPLAN\s+)?tests?,\s+(\d+)\s+failures\s*$", output)
    all_passed = re.findall(r"(?m)^PASS\s+(\S+)\s*$", output)
    passed = all_passed if expected_names is None else [name for name in all_passed if name in expected_names]
    if len(summaries) != 1:
        raise ValueError("Exactly one executed C# test summary is required.")
    total, failures = map(int, summaries[0])
    if expected_names is not None and (len(expected_names) != total or set(passed) != set(expected_names)):
        raise ValueError("Registered C# groups must each have an actual successful execution row.")
    if total <= 0 or failures or len(passed) != total or len(set(passed)) != total or re.search(r"(?m)^FAIL\s", output):
        raise ValueError("C# execution rows and successful summary do not reconcile.")
    return {"status": "passed", "total": total, "passed": len(passed), "failures": failures,
            "skippedTests": 0, "skippedGroups": 0, "passedTestNames": passed,
            "additionalPassedSubcaseNames": [name for name in all_passed if name not in passed]}


def parse_python_run(output: str) -> dict:
    summaries = re.findall(r"(?m)^Ran\s+(\d+)\s+tests?\s+in\s+", output)
    if len(summaries) != 1 or not re.search(r"(?m)^OK(?:\s*\(skipped=\d+\))?\s*$", output):
        raise ValueError("A successful verbose unittest execution summary is required.")
    passed = re.findall(r"(?m)^(test\S*\s+\(.+?\))\s+\.\.\.\s+ok\s*$", output)
    skipped = re.findall(r"(?m)^(.*?)\s+\.\.\.\s+skipped\s+.*$", output)
    skipped_groups = sum(name.startswith(("setUpClass (", "setUpModule (")) for name in skipped)
    skipped_tests = len(skipped) - skipped_groups
    total = int(summaries[0])
    summary_skips = re.search(r"(?m)^OK\s*\(skipped=(\d+)\)", output)
    if total <= 0 or total != len(passed) + skipped_tests or len(skipped) != (int(summary_skips[1]) if summary_skips else 0):
        raise ValueError("Python execution rows, skips and summary do not reconcile.")
    return {"status": "passed-with-skips" if skipped else "passed", "total": total,
            "passed": len(passed), "failures": 0, "skippedTests": skipped_tests,
            "skippedGroups": skipped_groups, "skippedNames": skipped}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def run_suite(command: list[str], label: str, parser, inputs: list[Path] | None = None) -> dict:
    hashes = {path.relative_to(ROOT).as_posix(): sha256(path) for path in inputs or []}
    started = utc_now()
    output = run_checked(command)
    result = parser(output)
    for relative, digest in hashes.items():
        if sha256(ROOT / relative) != digest:
            raise RuntimeError("A test input changed during evidence collection.")
    result.update({"command": label, "startedUtc": started, "completedUtc": utc_now(),
                   "exitCode": 0, "outputSha256": hashlib.sha256(output.encode("utf-8")).hexdigest(),
                   "inputSha256": hashes})
    print(f"Executed {label}: {result['passed']} passed; {result['skippedTests']} skipped tests; {result['skippedGroups']} skipped groups.")
    return result


def workbook_inventory(path: Path) -> dict:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        counts = {}
        for name, key in [("Tables", "tables"), ("Constraints", "constraints"), ("Structures", "structures")]:
            if name not in workbook.sheetnames:
                raise ValueError("Distributed workbook is missing a data sheet.")
            counts[key] = sum(any(value is not None for value in row) for row in workbook[name].iter_rows(min_row=2, values_only=True))
        return {"path": path.relative_to(ROOT).as_posix(), "sha256": sha256(path),
                "inspection": "Current distributed XLSX row inventory; not source-transcription or clinical validation.", **counts}
    finally:
        workbook.close()


def run_checked(command: list[str]) -> str:
    completed = subprocess.run(
        command,
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        timeout=900,
    )
    combined = "\n".join(
        part.strip() for part in (completed.stdout, completed.stderr) if part.strip()
    )
    if completed.returncode != 0:
        raise RuntimeError(
            "Evidence command failed ({}):\n{}".format(
                completed.returncode, combined
            )
        )
    return combined


def match_int(pattern: str, text: str, label: str) -> int:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    if not match:
        raise RuntimeError("Could not extract {} from command output.".format(label))
    return int(match.group(1))


def status_counts(rows: list[dict], field: str = "status") -> dict[str, int]:
    counts = Counter(str(row.get(field, "missing")).lower() for row in rows)
    return {key: counts[key] for key in sorted(counts)}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class EvidenceParserTests(unittest.TestCase):
    def test_csharp_requires_matching_executed_pass_rows(self):
        self.assertEqual(parse_csharp_run("PASS A\nPASS B\n2 tests, 0 failures")["passed"], 2)
        with self.assertRaises(ValueError):
            parse_csharp_run("PASS A\n2 tests, 0 failures")

    def test_registration_or_failed_output_is_not_execution_evidence(self):
        for output in ["376 registered tests", "PASS A\nFAIL B\n2 tests, 1 failures", "0 tests, 0 failures"]:
            with self.subTest(output=output), self.assertRaises(ValueError):
                parse_csharp_run(output)

    def test_dicom_summary_uses_its_actual_execution_rows(self):
        self.assertEqual(parse_csharp_run("PASS RTPLAN.A\n1 RTPLAN tests, 0 failures")["total"], 1)

    def test_nested_pass_rows_do_not_inflate_registered_group_count(self):
        result = parse_csharp_run("PASS nested_case\nPASS Group.A\n1 tests, 0 failures", ["Group.A"])
        self.assertEqual(result["passed"], 1)
        self.assertEqual(result["additionalPassedSubcaseNames"], ["nested_case"])
        with self.assertRaises(ValueError):
            parse_csharp_run("PASS Group.A\n1 tests, 0 failures", ["Group.A", "Group.B"])

    def test_python_class_skip_is_not_subtracted_from_executed_tests(self):
        output = "test_a (test_module.Parser) ... ok\nsetUpClass (test_module.Snapshot) ... skipped 'fixture absent'\nRan 1 test in 0.1s\nOK (skipped=1)"
        result = parse_python_run(output)
        self.assertEqual((result["total"], result["passed"], result["skippedTests"], result["skippedGroups"]), (1, 1, 0, 1))
        self.assertEqual(result["status"], "passed-with-skips")

    def test_python_case_skip_is_not_a_pass(self):
        output = "test_a (test_module.Parser) ... ok\ntest_b (test_module.Parser) ... skipped 'unsupported'\nRan 2 tests in 0.1s\nOK (skipped=1)"
        result = parse_python_run(output)
        self.assertEqual((result["total"], result["passed"], result["skippedTests"]), (2, 1, 1))

    def test_python_output_must_reconcile(self):
        for output in ["Ran 19 tests in 0.1s\nOK", "test_a (test_module.Parser) ... FAIL\nRan 1 test in 0.1s\nFAILED (failures=1)"]:
            with self.subTest(output=output), self.assertRaises(ValueError):
                parse_python_run(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--configuration", choices=["Debug", "Release"], default="Release")
    parser.add_argument("--release-version", default="3.2.0")
    parser.add_argument("--include-dicom", action="store_true")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    if args.self_test:
        result = unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(EvidenceParserTests))
        raise SystemExit(0 if result.wasSuccessful() else 1)
    scenario_paths = sorted(SCENARIO_DIR.glob("*.json"))
    scenarios = [json.loads(path.read_text(encoding="utf-8")) for path in scenario_paths]
    if not scenarios or not all(item.get("synthetic") is True for item in scenarios):
        raise RuntimeError("Every publication scenario must be explicitly synthetic.")

    mixed = json.loads(MIXED_SCENARIO.read_text(encoding="utf-8"))

    test_executable = (
        ROOT / "artifacts" / "bin" / args.configuration / "ClearPlan.Core.Tests.exe"
    )
    if not test_executable.is_file():
        raise RuntimeError(
            "Build ClearPlan.Core.Tests in the requested configuration before "
            "collecting paper evidence."
        )

    registered_names = re.findall(r'new\s+TestCase\s*\(\s*"([^"]+)"',
                                 (ROOT / "ClearPlan.Core.Tests/Program.cs").read_text(encoding="utf-8-sig"))
    csharp = run_suite([str(test_executable)], test_executable.relative_to(ROOT).as_posix(),
                       lambda text: parse_csharp_run(text, registered_names),
                       [test_executable, ROOT / "ClearPlan.Core.Tests/Program.cs", WORKBOOK,
                        ROOT / "ClearPlan.Script/Distribution/ConstraintTemplates/ClearPlan_StockConstraints2024.xlsx"] +
                       sorted(test_executable.parent.glob("ClearPlan.*.dll")))
    csharp["configuration"] = args.configuration
    raystation = run_suite([sys.executable, "-m", "unittest", "discover", "-s", "examples/raystation/tests", "-v"],
                          "python -m unittest discover -s examples/raystation/tests -v", parse_python_run,
                          sorted((ROOT / "examples/raystation").rglob("*.py")))
    raystation["runtimeValidation"] = False
    stock = run_suite([sys.executable, "tools/stock-catalog/test_stock_catalog.py", "-v"],
                     "python tools/stock-catalog/test_stock_catalog.py -v", parse_python_run,
                     [ROOT / "tools/stock-catalog" / name for name in
                      ["test_stock_catalog.py", "extract_stock_catalog.py", "stock2024.mapping.json"]])
    stock["sourceFixtureAvailable"] = (ROOT / "artifacts/stock-catalog/stock2024.normalized.json").is_file()
    stock["scope"] = "Parser tests always run. Source-transcription group is explicitly skipped when its private intermediate is absent; no skipped test is counted as passed."
    dicom = {"status": "not-run", "reason": "Use --include-dicom after a locked synthetic DICOM build."}
    if args.include_dicom:
        executable = ROOT / "artifacts" / "dicom-tests" / args.configuration / "ClearPlan.Dicom.Tests.exe"
        if not executable.is_file():
            raise RuntimeError("Build the synthetic DICOM tests before --include-dicom.")
        dicom = run_suite([str(executable)], executable.relative_to(ROOT).as_posix(), parse_csharp_run,
                          [executable] + sorted(executable.parent.glob("ClearPlan.*.dll")))

    workbook_output = run_checked(
        [sys.executable, str(ROOT / "tools" / "validate-constraint-workbook.py"), str(WORKBOOK)]
    )
    workbook_tables = match_int(r"tables=(\d+)", workbook_output, "workbook table count")
    workbook_constraints = match_int(
        r"constraints=(\d+)", workbook_output, "workbook constraint count"
    )
    workbook_structures = match_int(
        r"structures=(\d+)", workbook_output, "workbook structure count"
    )

    workbooks = [workbook_inventory(ROOT / "ClearPlan.Script/Distribution/ConstraintTemplates/ClearPlan_StockConstraints2024.xlsx"),
                 workbook_inventory(WORKBOOK)]
    if (workbooks[1]["tables"], workbooks[1]["constraints"], workbooks[1]["structures"]) != (workbook_tables, workbook_constraints, workbook_structures):
        raise RuntimeError("Default workbook validator and current inventory disagree.")
    workbooks[0]["loaderRegressionExecuted"] = "ExcelConstraintSource.Stock2024" in csharp["passedTestNames"]
    if not workbooks[0]["loaderRegressionExecuted"]:
        raise RuntimeError("The current C# run did not execute the distributed stock workbook regression.")
    publication_source = ROOT / "ClearPlan.Core/Simulation/SyntheticPublicationScenarioFactory.cs"
    publication_ids = sorted(set(re.findall(r'"(publication-(?:single|dual)-layer)"', publication_source.read_text(encoding="utf-8-sig"))))
    publication_tests = [name for name in csharp["passedTestNames"] if name.startswith("SyntheticPublication")]
    if len(publication_ids) != 2 or not publication_tests:
        raise RuntimeError("Publication fixture source and current execution evidence are required.")
    if not re.fullmatch(r"\d+\.\d+\.\d+", args.release_version):
        raise ValueError("Release version must contain three numeric components.")
    evidence = {
        "schemaVersion": 2,
        "generatedUtc": utc_now(),
        "scope": (
            "Software verification and deterministic synthetic demonstration only; "
            "no patient or clinical outcome data."
        ),
        "repositoryRelease": "v" + args.release_version,
        "releaseStatus": "unpublished-candidate; this collector does not verify remote publication",
        "repositoryCommit": run_checked(["git", "rev-parse", "HEAD"]).strip(),
        "repositoryDirty": bool(run_checked(["git", "status", "--porcelain"]).strip()),
        "manuscriptReconciliationRequired": True,
        "publicationVerification": "Pending separate manuscript, figure, release-tag and asset verification. Executed software tests alone do not establish release readiness.",
        "softwareVerification": {
            "csharpTests": csharp,
            "illustrativeRayStationAdapterTests": raystation,
            "stockCatalogTests": stock,
            "syntheticDicomTests": dicom,
        },
        "constraintWorkbook": workbooks[0],
        "constraintWorkbooks": workbooks,
        "syntheticScenarios": {
            "count": len(scenarios) + len(publication_ids),
            "scenarioIds": [item["scenarioId"] for item in scenarios] + publication_ids,
            "allExplicitlySynthetic": True,
            "sourceFiles": [path.relative_to(ROOT).as_posix() for path in scenario_paths],
            "generatedPublicationFixtures": {"source": publication_source.relative_to(ROOT).as_posix(),
                "sourceSha256": sha256(publication_source), "scenarioIds": publication_ids,
                "executedContractTests": publication_tests,
                "limitation": "Shared analytical phantom; dose is not calculated from MLC apertures. Figure capture and visual QA are separate."},
        },
        "mixedReviewScenario": {
            "scenarioId": mixed["scenarioId"],
            "seed": mixed["seed"],
            "scenarioGeneratedUtc": mixed["generatedUtc"],
            "sourceCount": len(mixed["sources"]),
            "sourceStatusCounts": status_counts(mixed["sources"]),
            "planCount": len(mixed["plans"]),
            "pqm": {
                "rowCount": len(mixed["pqmRows"]),
                "statusCounts": status_counts(mixed["pqmRows"]),
                "achievedValues": [
                    {
                        "templateCode": row["templateCode"],
                        "value": row["achievedValue"],
                        "unit": row["unit"],
                        "status": row["status"],
                    }
                    for row in mixed["pqmRows"]
                ],
            },
            "planCheck": {
                "rowCount": len(mixed["planCheckRows"]),
                "statusCounts": status_counts(mixed["planCheckRows"]),
                "injectedNonPassRows": sum(
                    1
                    for row in mixed["planCheckRows"]
                    if row["status"] != "pass"
                    and row["message"].startswith("Injected synthetic finding:")
                ),
            },
            "fields": {
                "rowCount": len(mixed["fieldRows"]),
                "idStatusCounts": status_counts(mixed["fieldRows"], "idStatus"),
                "nameStatusCounts": status_counts(mixed["fieldRows"], "nameStatus"),
                "suggestedNames": [row["suggestedName"] for row in mixed["fieldRows"]],
            },
            "mappings": {
                "rowCount": len(mixed["structureMappings"]),
                "statusCounts": status_counts(mixed["structureMappings"]),
            },
            "dvh": {
                "seriesCount": len(mixed["dvhSeries"]),
                "selectedSeriesCount": sum(
                    1 for series in mixed["dvhSeries"] if series["selected"]
                ),
                "pointsPerSeries": {
                    series["stableId"]: len(series["points"])
                    for series in mixed["dvhSeries"]
                },
            },
        },
        "figureInputs": {
            "status": "not-verified-by-this-collector",
            "permittedSources": "Explicitly synthetic simulator fixtures and original architecture diagrams only.",
        },
    }

    output_path = args.output if args.output.is_absolute() else ROOT / args.output
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(evidence, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print("Wrote {}".format(output_path))


if __name__ == "__main__":
    main()
