"""Portable structural tests; visual page inspection remains a separate gate."""
import importlib.util
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

spec = importlib.util.spec_from_file_location("builder", Path(__file__).with_name("build-technical-note-docx.py"))
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)

class ManuscriptBuilderTests(unittest.TestCase):
    def test_comment_anchors_not_printed(self):
        with tempfile.TemporaryDirectory(prefix="ClearPlan-docx-test-") as directory:
            source = Path(directory) / "input.md"
            source.write_text("# ClearPlan\n\nA claim [1].<!--ref:example--><!--anchor:section:Abstract-->\n", encoding="utf-8")
            doc = Document()
            builder.configure_styles(doc)
            builder.render_markdown(doc, source)
            output = Path(directory) / "output.docx"
            doc.save(output)
            reopened = Document(output)
            self.assertEqual(reopened.paragraphs[-1].text, "A claim [1].")
            self.assertNotIn("ref:example", reopened._element.xml)

    def test_table_has_repeating_header_borders_and_no_fixed_heights(self):
        doc = Document()
        builder.configure_styles(doc)
        builder.add_table(doc, ["Metric", "Definition", "Scope"], [["CI", "Formula", "Whole plan"]], ["---"] * 3)
        xml = doc.tables[0]._tbl.xml
        self.assertIn("tblHeader", xml)
        self.assertIn('w:color="D9D9D9"', xml)
        self.assertNotIn("trHeight", xml)

    def test_rate_equation_is_editable_omml(self):
        doc = Document()
        builder.add_rate_equation(doc)
        self.assertEqual(len(doc._element.findall(".//" + qn("m:f"))), 3)
        self.assertGreater(len(doc._element.findall(".//" + qn("m:sSub"))), 4)

    def test_headings_are_black(self):
        doc = Document()
        builder.configure_styles(doc)
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            self.assertEqual(str(doc.styles[name].font.color.rgb), "000000")

    def test_table_header_stays_with_first_data_row_only(self):
        doc = Document()
        builder.configure_styles(doc)
        builder.add_table(doc, ["Metric", "Definition", "Scope"],
                          [["CI", "Formula", "Plan"], ["GI", "Ratio", "Plan"]], ["---"] * 3)
        header = doc.tables[0].rows[0]
        self.assertTrue(all(p.paragraph_format.keep_with_next is True for c in header.cells for p in c.paragraphs))
        self.assertTrue(all(p.paragraph_format.keep_with_next is not True
                            for c in doc.tables[0].rows[-1].cells for p in c.paragraphs))

    def test_table_caption_stays_with_header(self):
        with tempfile.TemporaryDirectory(prefix="ClearPlan-docx-test-") as directory:
            source = Path(directory) / "table.md"
            source.write_text("**Table 1. Definitions.**\n\n| Metric | Value |\n| --- | --- |\n| CI | 1 |\n", encoding="utf-8")
            doc = Document()
            builder.configure_styles(doc)
            builder.render_markdown(doc, source)
            caption = doc.paragraphs[0]
            self.assertTrue(caption.paragraph_format.keep_with_next)
            self.assertTrue(caption.paragraph_format.keep_together)

    def test_figure_caption_lines_stay_together(self):
        doc = Document()
        builder.configure_styles(doc)
        self.assertTrue(doc.styles["Figure Caption"].paragraph_format.keep_together)
        self.assertFalse(doc.styles["Figure Caption"].paragraph_format.keep_with_next)

    def test_tall_figure_reserves_space_for_its_caption(self):
        with tempfile.TemporaryDirectory(prefix="ClearPlan-docx-test-") as directory:
            image = Path(directory) / "tall-synthetic-test.png"
            Image.new("RGB", (800, 1200), "white").save(image)
            doc = Document()
            builder.set_cell_margins(doc.sections[0])
            builder.configure_styles(doc)
            builder.add_figure(doc, image)
            shape = doc.inline_shapes[0]
            self.assertLessEqual(shape.height, Inches(7.25))
            self.assertLessEqual(shape.width, Inches(6.45))
            self.assertAlmostEqual(shape.width / shape.height, 800 / 1200, places=5)
            self.assertTrue(doc.paragraphs[-1].paragraph_format.keep_with_next)
            self.assertTrue(doc.paragraphs[-1].paragraph_format.keep_together)

    def test_inline_emphasis_is_italic_without_literal_stars(self):
        doc = Document()
        paragraph = builder.add_paragraph(doc, "An *estimated plan trajectory*, not **measured** delivery.")
        self.assertEqual(paragraph.text, "An estimated plan trajectory, not measured delivery.")
        self.assertTrue(next(run for run in paragraph.runs if run.text == "estimated plan trajectory").italic)
        self.assertTrue(next(run for run in paragraph.runs if run.text == "measured").bold)
        product = builder.add_paragraph(doc, "x * y * z")
        self.assertEqual(product.text, "x * y * z")

    def test_short_math_subscripts_do_not_rewrite_identifiers(self):
        doc = Document()
        paragraph = builder.add_paragraph(doc, "TV_Rx² / V_50%Rx; m_i and D_p; PTV_60; `TV_Rx`.")
        self.assertEqual(paragraph.text, "TVRx² / V50%Rx; mi and Dp; PTV_60; TV_Rx.")
        self.assertEqual([run.text for run in paragraph.runs if run.font.subscript], ["Rx", "50%Rx", "i", "p"])

    def test_assembler_uses_two_panel_caption_without_editing_source(self):
        with tempfile.TemporaryDirectory(prefix="ClearPlan-docx-test-") as directory:
            root = Path(directory)
            body = root / "body.md"
            original = "# ClearPlan: configurable review\n\n**Figure 1. Architecture.** Original caption.\n\n**Figure 2. Synthetic review workspace.** Old (a), (b), and (c) panels from [FINAL_RELEASE].\n"
            body.write_text(original, encoding="utf-8")
            abstract = root / "abstract.md"
            abstract.write_text("# Abstract\n\nSynthetic test.\n", encoding="utf-8")
            for stem in ("Figure_1_ClearPlan_architecture", "Figure_2_ClearPlan_workspace"):
                Image.new("RGB", (8, 8), "white").save(root / (stem + ".png"))
            output = root / "assembled.md"
            subprocess.run([sys.executable, str(Path(__file__).with_name("assemble-technical-note.py")),
                            "--body", str(body), "--abstract", str(abstract), "--output", str(output),
                            "--figure-directory", str(root)], check=True, capture_output=True)
            rendered = output.read_text(encoding="utf-8")
            self.assertIn("(A) the nominal field setting", rendered)
            self.assertIn("(B) the first-control-point beam's-eye view", rendered)
            self.assertNotIn("(c)", rendered)
            self.assertEqual(rendered.count("[FINAL_RELEASE]"), 1)
            self.assertEqual(body.read_text(encoding="utf-8"), original)

    def test_title_has_no_inherited_decorative_border(self):
        doc = Document()
        builder.configure_styles(doc)
        self.assertEqual(doc.styles["Title"]._element.findall(".//" + qn("w:pBdr")), [])

    def test_supplement_caption_is_unsplit_caption_text(self):
        with tempfile.TemporaryDirectory(prefix="ClearPlan-docx-test-") as directory:
            source = Path(directory) / "supplement.md"
            source.write_text("**Supplementary Material S1. Synthetic report.** Full caption text.\n", encoding="utf-8")
            doc = Document()
            builder.configure_styles(doc)
            builder.render_markdown(doc, source)
            self.assertEqual(doc.paragraphs[0].style.name, "Figure Caption")
            self.assertTrue(doc.paragraphs[0].style.paragraph_format.keep_together)

    def test_reference_entry_does_not_split_its_doi_onto_another_page(self):
        with tempfile.TemporaryDirectory(prefix="ClearPlan-docx-test-") as directory:
            source = Path(directory) / "reference.md"
            source.write_text("12. Example author. Example title. https://doi.org/10.example/test\n", encoding="utf-8")
            doc = Document()
            builder.configure_styles(doc)
            builder.render_markdown(doc, source)
            self.assertTrue(doc.paragraphs[0].paragraph_format.keep_together)
            self.assertIsNot(doc.paragraphs[0].paragraph_format.keep_with_next, True)

    def test_declared_tolerance_exponent_is_a_superscript(self):
        paragraph = builder.add_paragraph(Document(), "A 10^−9 absolute tolerance.")
        self.assertEqual(paragraph.text, "A 10−9 absolute tolerance.")
        self.assertTrue(next(run for run in paragraph.runs if run.text == "−9").font.superscript)

if __name__ == "__main__":
    unittest.main(verbosity=2)
