#!/usr/bin/env python3
"""Compare normalized XLSX rows with their complete legacy CSV source rows."""

from __future__ import annotations

import argparse
import importlib.util
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


COMPARE_FIELDS = [
    "constraint_id",
    "table_id",
    "structure_id",
    "structure_name",
    "metric",
    "unit",
    "comparator",
    "goal",
    "variation",
    "priority",
    "source",
    "comment",
    "legacy_metadata",
]


def load_converter(script_path: Path):
    spec = importlib.util.spec_from_file_location(
        "clearplan_constraint_converter", script_path
    )
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load converter module: {script_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def worksheet_records(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Constraints"]
        values = list(worksheet.iter_rows(values_only=True))
        headers = [str(value or "") for value in values[0]]
        return [
            dict(zip(headers, row, strict=False))
            for row in values[1:]
            if any(value not in (None, "") for value in row)
        ]
    finally:
        workbook.close()


def normalized(value: Any) -> Any:
    if value in ("", None):
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", action="append", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    converter = load_converter(root / "tools" / "convert-constraint-csvs-to-xlsx.py")
    inputs = [path.resolve() for path in args.input]
    _, expected_rows, _ = converter.convert(inputs)
    actual_rows = worksheet_records(args.workbook.resolve())
    if len(expected_rows) != len(actual_rows):
        raise AssertionError(
            f"row count differs: expected={len(expected_rows)} actual={len(actual_rows)}"
        )
    for index, (expected, actual) in enumerate(
        zip(expected_rows, actual_rows, strict=True), start=2
    ):
        for field in COMPARE_FIELDS:
            if normalized(expected.get(field)) != normalized(actual.get(field)):
                raise AssertionError(
                    f"Constraints!{index} field {field}: "
                    f"expected={expected.get(field)!r} actual={actual.get(field)!r}"
                )
        raw = json.loads(str(actual["legacy_metadata"]))
        source_path = inputs[
            [path.stem for path in inputs].index(str(actual["table_id"]))
        ]
        source_rows = converter.read_csv(source_path)
        source_index = int(str(actual["constraint_id"]).rsplit("-", 1)[1]) - 1
        if raw != source_rows[source_index]:
            raise AssertionError(
                f"Constraints!{index}: complete legacy row metadata differs"
            )
    counts = Counter(str(row["table_id"]) for row in actual_rows)
    for table_id in sorted(counts, key=str.casefold):
        print(f"{table_id}: {counts[table_id]} rows equivalent")
    print(f"PASS total={len(actual_rows)} rows equivalent")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
