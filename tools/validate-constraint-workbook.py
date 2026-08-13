#!/usr/bin/env python3
"""Validate normalized ClearPlan constraint workbooks and Unicode round-trips."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_HEADERS = {
    "Tables": {"table_id", "display_name", "active", "is_plan_sum"},
    "Constraints": {
        "constraint_id",
        "table_id",
        "structure_id",
        "metric",
        "unit",
        "comparator",
        "goal",
        "variation",
    },
    "Structures": {"structure_id", "canonical_name", "aliases"},
}

EXPECTED_UNICODE = set("äöüÄÖÜß≤≥³")


def normalized_id(value: object) -> str:
    return str(value or "").strip().casefold()


def duplicate_ids(values: list[str]) -> list[str]:
    seen: dict[str, str] = {}
    duplicates: set[str] = set()
    for value in values:
        key = normalized_id(value)
        if key in seen:
            duplicates.add(f"{seen[key]!r} / {value!r}")
        else:
            seen[key] = value
    return sorted(duplicates, key=str.casefold)


def sheet_records(worksheet) -> tuple[list[str], list[dict[str, object]]]:
    values = list(worksheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"{worksheet.title}: sheet is empty")
    headers = [str(value or "").strip() for value in values[0]]
    records = [
        dict(zip(headers, row, strict=False))
        for row in values[1:]
        if any(value not in (None, "") for value in row)
    ]
    return headers, records


def validate(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = (set(REQUIRED_HEADERS) | {"README"}) - set(
            workbook.sheetnames
        )
        if missing_sheets:
            raise ValueError(f"missing sheets: {sorted(missing_sheets)}")
        records: dict[str, list[dict[str, object]]] = {}
        all_text: list[str] = []
        for sheet_name, required in REQUIRED_HEADERS.items():
            headers, sheet_rows = sheet_records(workbook[sheet_name])
            missing_headers = required - set(headers)
            if missing_headers:
                raise ValueError(
                    f"{sheet_name}: missing headers {sorted(missing_headers)}"
                )
            records[sheet_name] = sheet_rows
            all_text.extend(
                str(value)
                for row in sheet_rows
                for value in row.values()
                if value is not None
            )
        all_text.extend(
            str(cell.value)
            for row in workbook["README"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        table_ids = [str(row["table_id"]) for row in records["Tables"]]
        constraint_ids = [
            str(row["constraint_id"]) for row in records["Constraints"]
        ]
        structure_ids = [
            str(row["structure_id"]) for row in records["Structures"]
        ]
        for id_name, values in (
            ("table_id", table_ids),
            ("constraint_id", constraint_ids),
            ("structure_id", structure_ids),
        ):
            duplicates = duplicate_ids(values)
            if duplicates:
                raise ValueError(
                    f"duplicate {id_name} (case-insensitive): {duplicates}"
                )
        known_tables = {normalized_id(value) for value in table_ids}
        unknown_tables = {
            str(row["table_id"])
            for row in records["Constraints"]
            if normalized_id(row["table_id"]) not in known_tables
        }
        if unknown_tables:
            raise ValueError(
                f"constraints reference unknown tables: {sorted(unknown_tables)}"
            )
        known_structures = {normalized_id(value) for value in structure_ids}
        unknown_structures = {
            str(row["structure_id"])
            for row in records["Constraints"]
            if normalized_id(row["structure_id"]) not in known_structures
        }
        if unknown_structures:
            raise ValueError(
                "constraints reference unknown structures: "
                f"{sorted(unknown_structures)}"
            )
        combined = "\n".join(all_text)
        missing_unicode = EXPECTED_UNICODE - set(combined)
        if missing_unicode:
            raise ValueError(
                f"Unicode sentinel characters missing: {sorted(missing_unicode)}"
            )
        if "?" in combined:
            raise ValueError("unexpected replacement question mark in workbook text")
        return {
            "tables": len(table_ids),
            "constraints": len(constraint_ids),
            "structures": len(structure_ids),
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbooks", nargs="*", type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    workbooks = args.workbooks or [
        root
        / "ClearPlan.Script"
        / "Distribution"
        / "ConstraintTemplates"
        / "ClearPlan_DefaultConstraints.xlsx"
    ]
    for workbook in workbooks:
        if not workbook.is_file():
            print(f"Constraint workbook validation skipped; missing: {workbook}")
            continue
        counts = validate(workbook)
        print(
            f"PASS {workbook}: tables={counts['tables']} "
            f"constraints={counts['constraints']} structures={counts['structures']}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
